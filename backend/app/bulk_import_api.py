from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook, Workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from .database import get_db
from .models import EmployeeMaster, MasterData, MasterType
from .admin_api import (
    get_current_user,
    require_permission,
    audit,
    now_local,
    create_password_hash,
    apply_user_master_refs,
    resolve_department,
    resolve_role,
)


router = APIRouter(
    prefix="/api/v1/admin/bulk-import",
    tags=["Dorado CRM Bulk Import"],
)

SUPPORTED_IMPORT_TYPES = {
    "MASTER_DATA",
    "EMPLOYEE",
    "LEAD",
    "CUSTOMER",
}

MASTER_CODE_ALIASES = {
    "code",
    "master_code",
    "id",
    "master_id",
    "type_id",
}

MASTER_NAME_ALIASES = {
    "name",
    "master_name",
    "value",
    "description",
    "title",
}

DISPLAY_ORDER_ALIASES = {
    "display_order",
    "order",
    "sort_order",
    "sequence",
    "sl_no",
    "serial",
}

ACTIVE_ALIASES = {
    "active",
    "is_active",
    "status",
}

DEPARTMENT_ALIASES = {
    "department_id",
    "department",
}


def clean_value(value: Any):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def normalize_header(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip().lower()

    for old, new in {
        " ": "_",
        "-": "_",
        "/": "_",
        ".": "",
        "(": "",
        ")": "",
        "&": "and",
    }.items():
        text = text.replace(old, new)

    while "__" in text:
        text = text.replace("__", "_")

    return text.strip("_")


def first_value(row: dict[str, Any], aliases: set[str]):
    for key in aliases:
        value = row.get(key)
        if value is not None and str(value).strip():
            return value
    return None


def parse_bool(value: Any, default=True) -> bool:
    if value is None:
        return default

    if isinstance(value, bool):
        return value

    text = str(value).strip().lower()

    if text in {"true", "yes", "y", "1", "active"}:
        return True

    if text in {"false", "no", "n", "0", "inactive"}:
        return False

    return default


def parse_int(value: Any, default=0) -> int:
    if value is None:
        return default

    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def workbook_from_bytes(content: bytes):
    try:
        return load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid Excel file: {exc}",
        )


def read_sheet(ws, limit: int | None = None):
    rows = ws.iter_rows(values_only=True)

    try:
        first_row = next(rows)
    except StopIteration:
        return [], []

    headers = [
        normalize_header(value) or f"column_{index + 1}"
        for index, value in enumerate(first_row)
    ]

    output = []

    for excel_row_number, values in enumerate(rows, start=2):
        if limit is not None and len(output) >= limit:
            break

        row = {
            headers[index]: clean_value(value)
            for index, value in enumerate(values)
            if index < len(headers)
        }

        if not any(value is not None for value in row.values()):
            continue

        row["_excel_row"] = excel_row_number
        output.append(row)

    return headers, output


def get_workbook_sheet(
    content: bytes,
    sheet_name: str | None,
):
    wb = workbook_from_bytes(content)

    if not wb.sheetnames:
        raise HTTPException(
            status_code=400,
            detail="Workbook has no worksheets",
        )

    selected_sheet = sheet_name or wb.sheetnames[0]

    if selected_sheet not in wb.sheetnames:
        raise HTTPException(
            status_code=400,
            detail=f"Worksheet '{selected_sheet}' not found",
        )

    return wb, selected_sheet


def get_master_type(
    db: Session,
    master_type_id: str | None,
):
    type_id = (master_type_id or "").strip().upper()

    if not type_id:
        raise HTTPException(
            status_code=400,
            detail="Master Type is required",
        )

    row = (
        db.query(MasterType)
        .filter(MasterType.id == type_id)
        .first()
    )

    if row is None:
        raise HTTPException(
            status_code=404,
            detail="Master Type not found",
        )

    return row


def master_row_values(row: dict[str, Any]):
    code = first_value(row, MASTER_CODE_ALIASES)
    name = first_value(row, MASTER_NAME_ALIASES)
    display_order = first_value(row, DISPLAY_ORDER_ALIASES)
    active = first_value(row, ACTIVE_ALIASES)
    department = first_value(row, DEPARTMENT_ALIASES)

    code = str(code).strip().upper() if code is not None else ""
    name = str(name).strip() if name is not None else ""

    return {
        "code": code,
        "name": name,
        "display_order": parse_int(display_order, 0),
        "active": parse_bool(active, True),
        "department_id": (
            str(department).strip()
            if department is not None
            else None
        ),
    }


def validate_master_rows(
    db: Session,
    master_type_id: str,
    rows: list[dict[str, Any]],
):
    results = []
    seen_codes = set()

    for row in rows:
        values = master_row_values(row)
        errors = []

        if not values["code"]:
            errors.append("Code is required")

        if not values["name"]:
            errors.append("Name is required")

        code = values["code"]

        if code:
            if code in seen_codes:
                errors.append("Duplicate code inside uploaded file")
            else:
                seen_codes.add(code)

        existing = None

        if code:
            existing = (
                db.query(MasterData)
                .filter(
                    MasterData.master_type_id == master_type_id,
                    func.upper(MasterData.code) == code.upper(),
                )
                .first()
            )

        action = "ERROR"

        if not errors:
            action = "UPDATE" if existing else "CREATE"

        results.append({
            "excel_row": row.get("_excel_row"),
            **values,
            "action": action,
            "valid": not errors,
            "errors": errors,
        })

    return results




# ============================================================
# EMPLOYEE BULK IMPORT HELPERS
# ============================================================

EMPLOYEE_TEXT_FIELDS = {
    "employee_name",
    "father_name",
    "gender",
    "mobile",
    "alternate_mobile",
    "email",
    "designation",
    "employee_type",
    "work_location",
    "blood_group",
    "emergency_contact_name",
    "emergency_contact_mobile",
    "emergency_contact_relation",
    "uan_number",
    "pf_number",
    "esic_number",
    "bank_name",
    "bank_account_holder_name",
    "bank_account_number",
    "bank_ifsc",
    "bank_branch",
    "pan_number",
    "aadhaar_number",
}

EMPLOYEE_BOOL_FIELDS = {
    "active",
    "pf_applicable",
    "esic_applicable",
    "force_password_reset",
}

EMPLOYEE_FLOAT_FIELDS = {
    "basic_salary",
    "hra",
    "conveyance_allowance",
    "other_allowance",
    "gross_salary",
    "ctc",
    "employee_pf_contribution",
    "employer_pf_contribution",
    "employee_esic_contribution",
    "employer_esic_contribution",
}

EMPLOYEE_DATE_FIELDS = {
    "date_of_birth",
    "date_of_joining",
    "date_of_exit",
}


def parse_float(value: Any):
    value = clean_value(value)

    if value is None:
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        raise ValueError(f"Invalid number: {value}")


def parse_employee_date(value: Any):
    value = clean_value(value)

    if value is None:
        return None

    # openpyxl normally returns Excel date cells as datetime/date.
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return value

    text = str(value).strip()

    from datetime import datetime

    formats = (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
    )

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    raise ValueError(
        f"Invalid date: {text}. Use DD-MM-YYYY or YYYY-MM-DD"
    )


def employee_row_values(row: dict[str, Any]):
    employee_id = clean_value(
        row.get("employee_id")
        or row.get("employee_code")
        or row.get("id")
    )

    if employee_id is not None:
        employee_id = str(employee_id).strip().upper()

    values: dict[str, Any] = {
        "employee_id": employee_id,
    }

    for field in EMPLOYEE_TEXT_FIELDS:
        values[field] = clean_value(row.get(field))

    for field in EMPLOYEE_BOOL_FIELDS:
        raw = clean_value(row.get(field))

        if raw is None:
            values[field] = None
        else:
            default = True if field in {"active", "force_password_reset"} else False
            values[field] = parse_bool(raw, default=default)

    for field in EMPLOYEE_FLOAT_FIELDS:
        raw = clean_value(row.get(field))
        values[field] = parse_float(raw) if raw is not None else None

    for field in EMPLOYEE_DATE_FIELDS:
        raw = clean_value(row.get(field))
        values[field] = parse_employee_date(raw) if raw is not None else None

    values["department_id"] = clean_value(
        row.get("department_id") or row.get("department")
    )

    values["role_id"] = clean_value(
        row.get("role_id") or row.get("app_role") or row.get("role")
    )

    values["reporting_manager_id"] = clean_value(
        row.get("reporting_manager_id")
        or row.get("reporting_manager")
    )

    values["temporary_password"] = clean_value(
        row.get("temporary_password")
        or row.get("password")
    )

    return values


def validate_employee_rows(
    db: Session,
    rows: list[dict[str, Any]],
):
    results = []
    seen_ids: set[str] = set()

    for row in rows:
        errors = []

        try:
            values = employee_row_values(row)
        except ValueError as exc:
            values = {
                "employee_id": clean_value(
                    row.get("employee_id")
                    or row.get("employee_code")
                    or row.get("id")
                ),
                "employee_name": clean_value(row.get("employee_name")),
            }
            errors.append(str(exc))

        employee_id = str(
            values.get("employee_id") or ""
        ).strip().upper()

        employee_name = str(
            values.get("employee_name") or ""
        ).strip()

        if not employee_id:
            errors.append("Employee ID is required")

        if not employee_name:
            errors.append("Employee Name is required")

        if employee_id:
            if employee_id in seen_ids:
                errors.append(
                    f"Duplicate Employee ID inside Excel: {employee_id}"
                )
            else:
                seen_ids.add(employee_id)

        existing = None

        if employee_id:
            existing = (
                db.query(EmployeeMaster)
                .filter(EmployeeMaster.id == employee_id)
                .first()
            )

        password = str(
            values.get("temporary_password") or ""
        )

        if existing is None:
            if len(password) < 6:
                errors.append(
                    "Temporary Password is required for new employee "
                    "and must be at least 6 characters"
                )
        elif password and len(password) < 6:
            errors.append(
                "Temporary Password must be at least 6 characters"
            )

        department_value = values.get("department_id")

        if department_value:
            if resolve_department(db, str(department_value)) is None:
                errors.append(
                    f"Invalid Department: {department_value}"
                )

        role_value = values.get("role_id")

        if role_value:
            if resolve_role(db, str(role_value)) is None:
                errors.append(
                    f"Invalid Role: {role_value}"
                )

        manager_value = values.get("reporting_manager_id")

        if manager_value:
            manager_text = str(manager_value).strip()

            manager = (
                db.query(EmployeeMaster)
                .filter(
                    (EmployeeMaster.id == manager_text)
                    | (EmployeeMaster.employee_name == manager_text)
                )
                .first()
            )

            if manager is None:
                errors.append(
                    f"Invalid Reporting Manager: {manager_text}"
                )

        results.append({
            "excel_row": row.get("_excel_row"),
            "employee_id": employee_id or None,
            "employee_name": employee_name or None,
            "action": "UPDATE" if existing else "CREATE",
            "valid": not errors,
            "errors": errors,
        })

    return results


def apply_employee_import_values(
    db: Session,
    employee: EmployeeMaster,
    values: dict[str, Any],
    *,
    creating: bool,
):
    for field in EMPLOYEE_TEXT_FIELDS:
        value = values.get(field)

        # On update, blank Excel cells do not erase existing data.
        if value is not None:
            setattr(employee, field, value)

    for field in EMPLOYEE_BOOL_FIELDS:
        value = values.get(field)

        if value is not None:
            setattr(employee, field, value)

    for field in EMPLOYEE_FLOAT_FIELDS:
        value = values.get(field)

        if value is not None:
            setattr(employee, field, value)

    for field in EMPLOYEE_DATE_FIELDS:
        value = values.get(field)

        if value is not None:
            setattr(employee, field, value)

    ref_payload = {}

    for field in (
        "department_id",
        "role_id",
        "reporting_manager_id",
    ):
        value = values.get(field)

        if value is not None:
            ref_payload[field] = value

    if ref_payload:
        apply_user_master_refs(
            db,
            employee,
            ref_payload,
        )

    password = str(
        values.get("temporary_password") or ""
    )

    if creating:
        employee.password = create_password_hash(password)

        if values.get("force_password_reset") is None:
            employee.force_password_reset = True
    elif password:
        employee.password = create_password_hash(password)



@router.get("/types")
def bulk_import_types(
    user: EmployeeMaster = Depends(get_current_user),
):
    return {
        "import_types": [
            {"id": "MASTER_DATA", "name": "Master Data"},
            {"id": "EMPLOYEE", "name": "Employee / User"},
            {"id": "LEAD", "name": "Lead"},
            {"id": "CUSTOMER", "name": "Customer"},
        ]
    }


@router.get("/master-types")
def bulk_master_types(
    user: EmployeeMaster = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = (
        db.query(MasterType)
        .filter(MasterType.active.is_(True))
        .order_by(
            MasterType.display_order,
            MasterType.master_name,
        )
        .all()
    )

    return [
        {
            "id": row.id,
            "name": row.master_name,
        }
        for row in rows
    ]



@router.get("/template")
def download_bulk_import_template(
    import_type: str,
    master_type_id: str | None = None,
    user: EmployeeMaster = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    import_type = import_type.strip().upper()

    if import_type == "EMPLOYEE":
        require_permission(db, user, "user_management", "view")
    else:
        require_permission(db, user, "data_master", "view")

    if import_type not in SUPPORTED_IMPORT_TYPES:
        raise HTTPException(
            status_code=400,
            detail="Unsupported import type",
        )

    wb = Workbook()
    ws = wb.active

    if import_type == "MASTER_DATA":
        master_type = get_master_type(db, master_type_id)

        safe_name = (
            master_type.master_name
            .replace("/", "_")
            .replace("\\", "_")
            .replace(" ", "_")
        )

        ws.title = safe_name[:31]

        headers = [
            "code",
            "name",
            "display_order",
            "active",
            "department_id",
        ]

        filename = f"{safe_name}_Import_Template.xlsx"

    elif import_type == "EMPLOYEE":
        ws.title = "Employee"

        headers = [
            "employee_id",
            "employee_name",
            "father_name",
            "date_of_birth",
            "gender",
            "mobile",
            "alternate_mobile",
            "email",
            "designation",
            "department_id",
            "role_id",
            "reporting_manager_id",
            "date_of_joining",
            "employee_type",
            "work_location",
            "active",
            "blood_group",
            "emergency_contact_name",
            "emergency_contact_mobile",
            "emergency_contact_relation",
            "pf_applicable",
            "esic_applicable",
            "uan_number",
            "pf_number",
            "esic_number",
            "basic_salary",
            "hra",
            "conveyance_allowance",
            "other_allowance",
            "gross_salary",
            "ctc",
            "bank_name",
            "bank_account_holder_name",
            "bank_account_number",
            "bank_ifsc",
            "bank_branch",
            "pan_number",
            "aadhaar_number",
            "temporary_password",
            "force_password_reset",
        ]

        filename = "Employee_User_Import_Template.xlsx"

    elif import_type == "LEAD":
        ws.title = "Lead"

        headers = [
            "lead_date",
            "lead_id",
            "executive",
            "lead_source",
            "customer_name",
            "mobile",
            "location",
            "building_type",
            "lift_type",
            "requirement_time",
            "call_done",
            "appointment_fixed",
            "appointment_date",
            "site_visit_done",
            "site_visit_date",
            "survey_done",
            "survey_date",
            "quotation_given",
            "quotation_date",
            "negotiation_status",
            "negotiation_date",
            "order_finalized",
            "order_date",
            "order_lost",
            "lost_date",
            "lead_status",
            "remarks",
            "follow_up_date",
            "assigned_marketing_head",
            "lead_priority",
            "expected_order_value",
            "expected_close_date",
            "lost_reason",
            "active",
        ]

        filename = "Lead_Import_Template.xlsx"

    else:
        ws.title = "Customer"

        headers = [
            "customer_id",
            "customer_name",
            "mobile",
            "alternate_mobile",
            "email",
            "address",
            "location",
            "building_type",
            "contact_person",
            "remarks",
            "active",
        ]

        filename = "Customer_Import_Template.xlsx"

    # Header row
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = cell.font.copy(bold=True)

        width = max(15, min(len(header) + 5, 30))
        ws.column_dimensions[cell.column_letter].width = width

    ws.freeze_panes = "A2"

    # Helpful example row only for Master Data
    if import_type == "MASTER_DATA":
        ws.cell(row=2, column=1, value="CODE001")
        ws.cell(row=2, column=2, value="Example - replace or delete this row")
        ws.cell(row=2, column=3, value=1)
        ws.cell(row=2, column=4, value="Yes")

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            ),
        },
    )

@router.post("/preview")
async def preview_bulk_import(
    import_type: str = Form(...),
    master_type_id: str | None = Form(None),
    sheet_name: str | None = Form(None),
    file: UploadFile = File(...),
    user: EmployeeMaster = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_permission(db, user, "data_master", "view")

    import_type = import_type.strip().upper()

    if import_type not in SUPPORTED_IMPORT_TYPES:
        raise HTTPException(
            status_code=400,
            detail="Unsupported import type",
        )

    filename = (file.filename or "").lower()

    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are supported",
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty",
        )

    wb, selected_sheet = get_workbook_sheet(
        content,
        sheet_name,
    )

    headers, rows = read_sheet(
        wb[selected_sheet],
        limit=100,
    )

    validation = None

    if import_type == "MASTER_DATA":
        master_type = get_master_type(
            db,
            master_type_id,
        )

        validation = validate_master_rows(
            db,
            master_type.id,
            rows,
        )

    elif import_type == "EMPLOYEE":
        validation = validate_employee_rows(
            db,
            rows,
        )

    return {
        "success": True,
        "filename": file.filename,
        "import_type": import_type,
        "master_type_id": (
            master_type_id.strip().upper()
            if master_type_id
            else None
        ),
        "available_sheets": wb.sheetnames,
        "selected_sheet": selected_sheet,
        "headers": headers,
        "rows": rows,
        "validation": validation,
        "preview_rows": len(rows),
    }


@router.post("/import")
async def execute_bulk_import(
    import_type: str = Form(...),
    master_type_id: str | None = Form(None),
    sheet_name: str | None = Form(None),
    update_existing: bool = Form(True),
    file: UploadFile = File(...),
    user: EmployeeMaster = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    import_type = import_type.strip().upper()

    if import_type == "EMPLOYEE":
        require_permission(db, user, "user_management", "add")
    else:
        require_permission(db, user, "data_master", "add")

    if import_type not in {"MASTER_DATA", "EMPLOYEE"}:
        raise HTTPException(
            status_code=400,
            detail=(
                "Database import for this type is not enabled yet. "
                "Use Preview first."
            ),
        )

    filename = (file.filename or "").lower()

    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are supported",
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty",
        )

    if import_type == "EMPLOYEE":
        wb, selected_sheet = get_workbook_sheet(
            content,
            sheet_name,
        )

        _, rows = read_sheet(wb[selected_sheet])

        validation = validate_employee_rows(
            db,
            rows,
        )

        validation_by_row = {
            item["excel_row"]: item
            for item in validation
        }

        created = 0
        updated = 0
        skipped = 0
        failed = 0
        results = []

        for raw_row in rows:
            excel_row = raw_row.get("_excel_row")

            item = validation_by_row.get(excel_row)

            if item is None:
                failed += 1
                results.append({
                    "excel_row": excel_row,
                    "action": "ERROR",
                    "valid": False,
                    "errors": ["Validation result not found"],
                })
                continue

            if not item["valid"]:
                failed += 1
                results.append(item)
                continue

            try:
                values = employee_row_values(raw_row)

                employee_id = str(
                    values.get("employee_id") or ""
                ).strip().upper()

                existing = (
                    db.query(EmployeeMaster)
                    .filter(EmployeeMaster.id == employee_id)
                    .first()
                )

                if existing and not update_existing:
                    skipped += 1

                    results.append({
                        **item,
                        "action": "SKIP",
                    })

                    continue

                # Per-row savepoint:
                # one bad employee will not roll back successful rows.
                with db.begin_nested():
                    if existing:
                        employee = existing

                        apply_employee_import_values(
                            db,
                            employee,
                            values,
                            creating=False,
                        )

                        action = "UPDATE"

                    else:
                        employee = EmployeeMaster(
                            id=employee_id,
                            employee_name=str(
                                values.get("employee_name") or ""
                            ).strip(),
                            active=(
                                values.get("active")
                                if values.get("active") is not None
                                else True
                            ),
                            date_of_joining=(
                                values.get("date_of_joining")
                                or now_local()
                            ),
                        )

                        apply_employee_import_values(
                            db,
                            employee,
                            values,
                            creating=True,
                        )

                        db.add(employee)

                        action = "CREATE"

                    db.flush()

                if action == "CREATE":
                    created += 1
                else:
                    updated += 1

                results.append({
                    **item,
                    "action": action,
                    "valid": True,
                    "errors": [],
                })

            except Exception as exc:
                failed += 1

                results.append({
                    **item,
                    "action": "ERROR",
                    "valid": False,
                    "errors": [str(exc)],
                })

        audit(
            db,
            user.id,
            "BULK_IMPORT",
            "user_management",
            "EmployeeMaster",
            None,
            (
                f"Employee import: "
                f"created={created}, "
                f"updated={updated}, "
                f"skipped={skipped}, "
                f"failed={failed}"
            ),
        )

        db.commit()

        return {
            "success": failed == 0,
            "filename": file.filename,
            "import_type": import_type,
            "selected_sheet": selected_sheet,
            "summary": {
                "total": len(rows),
                "created": created,
                "updated": updated,
                "skipped": skipped,
                "failed": failed,
            },
            "results": results,
        }

    master_type = get_master_type(
        db,
        master_type_id,
    )

    wb, selected_sheet = get_workbook_sheet(
        content,
        sheet_name,
    )

    _, rows = read_sheet(wb[selected_sheet])

    validation = validate_master_rows(
        db,
        master_type.id,
        rows,
    )

    created = 0
    updated = 0
    skipped = 0
    failed = 0
    results = []

    for item in validation:
        if not item["valid"]:
            failed += 1
            results.append(item)
            continue

        try:
            # Each Excel row gets its own database savepoint.
            # If this row fails, only this row is rolled back.
            with db.begin_nested():
                existing = (
                    db.query(MasterData)
                    .filter(
                        MasterData.master_type_id == master_type.id,
                        func.upper(MasterData.code)
                        == item["code"].upper(),
                    )
                    .first()
                )

                if existing:
                    if not update_existing:
                        skipped += 1
                        results.append({
                            **item,
                            "action": "SKIP",
                        })
                        continue

                    existing.name = item["name"]
                    existing.display_order = item["display_order"]
                    existing.active = item["active"]

                    if item["department_id"] is not None:
                        existing.department_id = item["department_id"]

                    if hasattr(existing, "updated_at"):
                        existing.updated_at = now_local()

                    updated += 1

                    results.append({
                        **item,
                        "action": "UPDATED",
                    })

                else:
                    new_row = MasterData(
                        master_type_id=master_type.id,
                        department_id=item["department_id"],
                        code=item["code"],
                        name=item["name"],
                        display_order=item["display_order"],
                        active=item["active"],
                    )

                    if hasattr(new_row, "created_at"):
                        new_row.created_at = now_local()

                    if hasattr(new_row, "updated_at"):
                        new_row.updated_at = now_local()

                    db.add(new_row)

                    created += 1

                    results.append({
                        **item,
                        "action": "CREATED",
                    })

        except Exception as exc:
            failed += 1

            results.append({
                **item,
                "action": "ERROR",
                "valid": False,
                "errors": [str(exc)],
            })

    try:
        audit(
            db,
            user.id,
            "BULK_IMPORT",
            "data_master",
            "MasterData",
            master_type.id,
            (
                f"File={file.filename}; "
                f"Sheet={selected_sheet}; "
                f"Created={created}; "
                f"Updated={updated}; "
                f"Skipped={skipped}; "
                f"Failed={failed}"
            ),
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "success": failed == 0,
        "import_type": import_type,
        "master_type_id": master_type.id,
        "master_type_name": master_type.master_name,
        "sheet": selected_sheet,
        "summary": {
            "total": len(validation),
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "failed": failed,
        },
        "results": results,
    }




